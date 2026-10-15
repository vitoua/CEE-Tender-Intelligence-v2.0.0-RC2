import os,json,csv,io,secrets,asyncio,re
from datetime import timedelta
from collections import Counter
from urllib.parse import urlencode
from fastapi import FastAPI,Request,Form,HTTPException
from fastapi.responses import HTMLResponse,JSONResponse,RedirectResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from app import core,engine,translation
from app.core import Tender,VERSION,BUILD,now,date,redact,safeurl
from app.registry import SOURCES,COUNTRIES
from app.i18n import tr
app=FastAPI(title='GOODRAM Tender Intelligence',version=VERSION)
app.mount('/static',StaticFiles(directory='app/static'),name='static')
tpl=Jinja2Templates(directory='app/templates')
def language(r):return r.cookies.get('lang','uk') if r.cookies.get('lang','uk') in ('uk','pl','en','es') else 'uk'
def context(r,**kw):return {'request':r,'t':tr(language(r)),'lang':language(r),'version':VERSION,'sources':SOURCES,'countries':COUNTRIES,'job':engine.JOB,'token_required':bool(os.getenv('ADMIN_TOKEN')),**kw}
def guard(r):
 token=os.getenv('ADMIN_TOKEN','')
 if token and not secrets.compare_digest(r.headers.get('X-Admin-Token',''),token):raise HTTPException(401,'Enter admin token in Settings')
 origin=r.headers.get('origin')
 if origin and origin.rstrip('/')!=str(r.base_url).rstrip('/'):raise HTTPException(403,'Cross-origin write rejected')
def select(r):
 q=r.query_params;state=q.get('state','active');country=q.get('country','');source=q.get('source','');term=q.get('q','').casefold();data=core.rows()
 counts=Counter('stale' if x.stale else x.state() for x in data)
 selected=[x for x in data if (state=='all' or (state=='stale' and x.stale) or (not x.stale and x.state()==state)) and (not country or x.country==country) and (not source or x.source==source) and (not term or term in (x.title+' '+x.buyer+' '+x.description).casefold())]
 selected.sort(key=lambda x:(x.deadline or '9999',x.title))
 return selected,counts
@app.get('/health')
def health():return {'status':'ok','version':VERSION,'build_id':BUILD,'records':len(core.rows()),'job':engine.JOB,'storage':'SQLite; persistent only if DATA_DIR is on persistent volume'}
@app.get('/lang/{code}')
def lang(code):
 r=RedirectResponse('/',303);r.set_cookie('lang',code if code in ('uk','pl','en','es') else 'uk',samesite='lax');return r
@app.get('/')
def home(request:Request):
 data,counts=select(request)
 try:page=max(1,int(request.query_params.get('page','1')))
 except ValueError:page=1
 total=len(data);data=data[(page-1)*50:page*50];qp=dict(request.query_params);qp.pop('page',None)
 return tpl.TemplateResponse(request=request,name='index.html',context=context(request,rows=data,counts=counts,total=total,page=page,pages=max(1,(total+49)//50),query=urlencode(qp),default_date=(now()-timedelta(days=30)).date().isoformat(),runs=core.runs()))
@app.post('/refresh/start')
async def start(request:Request,countries:list[str]=Form(default=[]),sources:list[str]=Form(default=[]),published_from:str=Form(''),min_value_eur:float=Form(0)):
 guard(request)
 if engine.JOB['running']:raise HTTPException(409,'Refresh already running')
 if not countries or not sources or any(x not in COUNTRIES for x in countries) or any(x not in SOURCES for x in sources):raise HTTPException(400,'Select valid countries and sources')
 start=date(published_from)
 if not start or start>now() or min_value_eur<0:raise HTTPException(400,'Invalid date or minimum value')
 if not any(SOURCES[s][1] in countries or (s=='ted' and any(c not in ('UKR','TUR') for c in countries)) for s in sources):raise HTTPException(400,'Selected sources do not cover selected countries')
 engine.JOB.update(running=True,progress=0,message='Queued')
 engine.TASK=asyncio.create_task(engine.execute(countries,sources,start,min_value_eur))
 return {'ok':True}
@app.get('/refresh/status')
def progress():return engine.JOB
@app.get('/tenders/{key}')
def detail(key,request:Request):
 x=next((x for x in core.rows() if x.key==key),None)
 if not x:raise HTTPException(404,'Tender not found')
 return tpl.TemplateResponse(request=request,name='detail.html',context=context(request,x=x,original=request.query_params.get('view')=='original'))
@app.get('/analytics')
def analytics(request:Request):
 data,counts=select(request);by_country=Counter(x.country for x in data);by_source=Counter(x.source for x in data);by_day=Counter(x.published[:10] for x in data if x.published)
 return tpl.TemplateResponse(request=request,name='analytics.html',context=context(request,counts=counts,total=len(data),by_country=by_country,by_source=by_source,by_day=dict(sorted(by_day.items())),max_count=max([1,*by_country.values(),*by_source.values(),*by_day.values()])))
@app.get('/cpv')
@app.get('/keywords')
def filters(request:Request):return tpl.TemplateResponse(request=request,name='filters.html',context=context(request,profiles=core.profiles(),mode='codes' if request.url.path=='/cpv' else 'keywords'))
@app.post('/filters/save')
async def save_filters(request:Request,source:str=Form(),mode:str=Form(),values:str=Form()):
 guard(request)
 if source not in SOURCES or mode not in ('codes','keywords'):raise HTTPException(400,'Invalid filter profile')
 vals=list(dict.fromkeys(x.strip() for x in values.splitlines() if x.strip()))
 if len(vals)>150 or any(len(x)>120 for x in vals):raise HTTPException(400,'Profile too large')
 if mode=='codes' and any(not re.fullmatch(r'\d{8}',x) for x in vals):raise HTTPException(400,'CPV must contain exactly 8 digits')
 p=core.profiles();p[source][mode]=vals;core.save_profiles(p);return {'ok':True,'message':'Saved. Refresh source to apply matching changes.'}
@app.get('/sources')
@app.get('/status')
def sources_page(request:Request):return tpl.TemplateResponse(request=request,name='status.html',context=context(request,runs=core.runs(),translator=translation.status(),limits=engine.LIMITS))
@app.get('/settings')
def settings(request:Request):return tpl.TemplateResponse(request=request,name='settings.html',context=context(request))
@app.get('/about')
def about(request:Request):return tpl.TemplateResponse(request=request,name='about.html',context=context(request))
@app.post('/translate/{key}')
async def translate_one(key,request:Request,provider:str=Form('free'),source_language:str=Form('auto')):
 guard(request);x=next((x for x in core.rows() if x.key==key),None)
 if not x:raise HTTPException(404)
 ok=await translation.translate(x,language(request),provider,source_language)
 return JSONResponse({'ok':ok,'translator':translation.status()},200 if ok else 400)
@app.post('/translate-visible')
async def translate_visible(request:Request,ids:list[str]=Form(default=[]),provider:str=Form('free'),source_language:str=Form('auto')):
 guard(request);wanted=set(ids[:20]);n=0
 for x in core.rows():
  if x.key in wanted:
   if not await translation.translate(x,language(request),provider,source_language):break
   n+=1
 return JSONResponse({'ok':n>0,'translated':n,'translator':translation.status()},200 if n else 400)
@app.post('/translator/test')
async def test_translator(request:Request,provider:str=Form('free')):
 guard(request);x=Tender('test','test','', 'Supply of enterprise SSD storage')
 ok=await translation.translate(x,language(request),provider,'en')
 return JSONResponse({'ok':ok,'result':x.translations,'translator':translation.status()},200 if ok else 400)
def cell(v):
 # Prevent spreadsheet formula injection in externally sourced text.
 return "'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v
@app.get('/export.{fmt}')
def export(fmt,request:Request):
 data,_=select(request);headers=['Source','Country','ID','Title','Buyer','Published UTC','Deadline UTC','Value','Currency','CPV','URL','State','Provenance','Notes','Stale'];records=[[x.source,x.country,x.id,x.title,x.buyer,x.published,x.deadline,x.value,x.currency,', '.join(x.cpv),x.url,x.state(),x.provenance,'; '.join(x.notes),x.stale] for x in data]
 if fmt=='json':return JSONResponse([x.to_dict()|{'state':x.state()} for x in data],headers={'Content-Disposition':'attachment; filename="GOODRAM-v2.json"'})
 if fmt=='csv':
  buf=io.StringIO();w=csv.writer(buf);w.writerow(headers);w.writerows([[cell(v) for v in row] for row in records]);return StreamingResponse(iter([('\ufeff'+buf.getvalue()).encode()]),media_type='text/csv',headers={'Content-Disposition':'attachment; filename="GOODRAM-v2.csv"'})
 if fmt!='xlsx':raise HTTPException(404)
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 wb=Workbook();ws=wb.active;ws.title='Tenders';ws.append(headers)
 for row in records:ws.append([cell(v) for v in row])
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='14283F')
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in col[:100])+2))
 buf=io.BytesIO();wb.save(buf);buf.seek(0)
 return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="GOODRAM-v2.xlsx"'})
